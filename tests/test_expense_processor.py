"""Tests for the expense processor."""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from processors.expense_processor import (
    ExpenseProcessorResult,
    process_expenses,
)
from models.expense import ExpenseCategory, ExpenseRecord, ExpenseStatus
from triggers.monthly_trigger import trigger_monthly_expense


class TestExpenseProcessor:
    """Tests for process_expenses function."""

    @pytest.fixture
    def expense_data(self, tmp_path):
        """Generate expense Excel and return path + client/month info."""
        run = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        excel_path = run.trigger_data["expense_file"]
        return excel_path, "TV-001", "2026-08"

    def test_returns_processor_result(self, expense_data):
        excel_path, client_id, month = expense_data
        result = process_expenses(excel_path, client_id, month)
        assert isinstance(result, ExpenseProcessorResult)

    def test_extracts_all_valid_records(self, expense_data):
        excel_path, client_id, month = expense_data
        result = process_expenses(excel_path, client_id, month)
        # 13 rows in the synthetic data
        assert result.success_count == 13
        assert result.error_count == 0

    def test_record_fields_are_populated(self, expense_data):
        excel_path, client_id, month = expense_data
        result = process_expenses(excel_path, client_id, month)
        first = result.records[0]
        assert first.expense_id == "EXP-0002"
        assert first.client_id == "TV-001"
        assert first.employee == "Alice Johnson"
        assert first.vendor == "Delta Airlines"
        assert first.amount == 487.00
        assert first.category == ExpenseCategory.TRAVEL
        assert first.month == "2026-08"
        assert first.status == ExpenseStatus.PENDING

    def test_records_have_expense_ids(self, expense_data):
        excel_path, client_id, month = expense_data
        result = process_expenses(excel_path, client_id, month)
        ids = [r.expense_id for r in result.records]
        # All IDs unique
        assert len(ids) == len(set(ids))

    def test_receipt_ref_populated_for_most_records(self, expense_data):
        excel_path, client_id, month = expense_data
        result = process_expenses(excel_path, client_id, month)
        # All 13 records have receipt references in the current test data
        with_ref = [r for r in result.records if r.receipt_reference is not None]
        assert len(with_ref) == 13

    def test_categories_are_valid_enums(self, expense_data):
        excel_path, client_id, month = expense_data
        result = process_expenses(excel_path, client_id, month)
        for record in result.records:
            assert isinstance(record.category, ExpenseCategory)

    def test_amounts_are_non_negative(self, expense_data):
        excel_path, client_id, month = expense_data
        result = process_expenses(excel_path, client_id, month)
        for record in result.records:
            assert record.amount >= 0

    def test_file_not_found_raises(self):
        with pytest.raises(FileNotFoundError):
            process_expenses("/nonexistent/path.xlsx", "TV-001", "2026-08")

    def test_invalid_excel_bad_amount(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Employee", "Date", "Description", "Vendor", "Amount", "Category", "Receipt Ref"])
        ws.append(["Test User", "2026-08-01", "Bad expense", "Test Vendor", "not_a_number", "Travel", ""])
        bad_path = tmp_path / "bad.xlsx"
        wb.save(str(bad_path))
        wb.close()
        result = process_expenses(bad_path, "TV-001", "2026-08")
        assert result.error_count == 1
        assert result.success_count == 0

    def test_invalid_excel_negative_amount(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Employee", "Date", "Description", "Vendor", "Amount", "Category", "Receipt Ref"])
        ws.append(["Test User", "2026-08-01", "Negative", "Test Vendor", -50, "Travel", ""])
        bad_path = tmp_path / "negative.xlsx"
        wb.save(str(bad_path))
        wb.close()
        result = process_expenses(bad_path, "TV-001", "2026-08")
        assert result.error_count == 1
        assert "Negative amount" in result.errors[0].error_message

    def test_invalid_excel_missing_employee(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Employee", "Date", "Description", "Vendor", "Amount", "Category", "Receipt Ref"])
        ws.append(["", "2026-08-01", "No employee", "Test Vendor", 100, "Travel", ""])
        bad_path = tmp_path / "no_emp.xlsx"
        wb.save(str(bad_path))
        wb.close()
        result = process_expenses(bad_path, "TV-001", "2026-08")
        assert result.error_count == 1
        assert "Missing employee" in result.errors[0].error_message

    def test_invalid_excel_unknown_category(self, tmp_path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Employee", "Date", "Description", "Vendor", "Amount", "Category", "Receipt Ref"])
        ws.append(["Test User", "2026-08-01", "Bad cat", "Test Vendor", 100, "Crypto", ""])
        bad_path = tmp_path / "bad_cat.xlsx"
        wb.save(str(bad_path))
        wb.close()
        result = process_expenses(bad_path, "TV-001", "2026-08")
        assert result.error_count == 1
        assert "Unrecognized category" in result.errors[0].error_message

    def test_result_to_dict(self, expense_data):
        excel_path, client_id, month = expense_data
        result = process_expenses(excel_path, client_id, month)
        d = result.to_dict()
        assert "records" in d
        assert "errors" in d
        assert d["success_count"] == 13
        assert d["error_count"] == 0
        assert isinstance(d["records"][0], dict)
