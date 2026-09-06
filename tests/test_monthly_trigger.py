"""Tests for the monthly expense trigger."""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from triggers.monthly_trigger import trigger_monthly_expense
from models.workflow import WorkflowRun, WorkflowStatus


class TestMonthlyTrigger:
    """Tests for trigger_monthly_expense function."""

    def test_creates_workflow_run(self, tmp_path):
        run = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        assert isinstance(run, WorkflowRun)
        assert run.client_id == "TV-001"
        assert run.month == "2026-08"
        assert run.status == WorkflowStatus.CREATED

    def test_run_has_unique_id(self, tmp_path):
        run1 = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        run2 = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        assert run1.run_id != run2.run_id

    def test_creates_expense_excel_file(self, tmp_path):
        run = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        excel_path = Path(run.trigger_data["expense_file"])
        assert excel_path.exists()
        assert excel_path.suffix == ".xlsx"

    def test_creates_receipt_json_file(self, tmp_path):
        run = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        receipt_path = Path(run.trigger_data["receipt_file"])
        assert receipt_path.exists()
        assert receipt_path.suffix == ".json"

    def test_trigger_data_contains_client_info(self, tmp_path):
        run = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        assert run.trigger_data["client_name"] == "TechVista Inc."
        assert run.trigger_data["contact"] == "Sarah Chen"
        assert run.trigger_data["industry"] == "Technology"

    def test_unknown_client_raises_error(self, tmp_path):
        with pytest.raises(ValueError, match="Unknown client_id"):
            trigger_monthly_expense("INVALID-999", "2026-08", output_dir=tmp_path)

    def test_excel_file_is_valid_workbook(self, tmp_path):
        run = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        from openpyxl import load_workbook
        wb = load_workbook(run.trigger_data["expense_file"], read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        # Header + 13 data rows
        assert len(rows) == 14
        assert rows[0][0] == "Employee"

    def test_receipt_json_is_valid(self, tmp_path):
        import json
        run = trigger_monthly_expense("TV-001", "2026-08", output_dir=tmp_path)
        with open(run.trigger_data["receipt_file"]) as f:
            data = json.load(f)
        assert isinstance(data, list)
        assert len(data) == 13
