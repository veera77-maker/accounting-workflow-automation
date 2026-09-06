"""Expense processor — extracts structured records from Excel spreadsheets.

This component handles the "Excel processing" step of the Case 3 workflow.
It reads a client's expense spreadsheet and produces validated ExpenseRecord objects.

Design decisions:
- Uses openpyxl for Excel reading (lightweight, no heavy dependencies).
- Validates each row against the ExpenseRecord schema.
- Collects malformed rows as errors rather than silently discarding them.
- Keeps processing pure (no side effects) — returns results, does not persist.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from pydantic import ValidationError

from models.expense import ExpenseCategory, ExpenseRecord, ExpenseStatus


# Column positions (0-indexed) expected in the expense spreadsheet.
# The trigger generates files with this exact header row.
EXPECTED_HEADERS = [
    "Employee", "Date", "Description", "Vendor",
    "Amount", "Category", "Receipt Ref",
]


class ExpenseProcessingError:
    """Records a single row-level processing failure."""

    def __init__(self, row_number: int, raw_row: dict, error_message: str):
        self.row_number = row_number
        self.raw_row = raw_row
        self.error_message = error_message

    def to_dict(self) -> dict:
        return {
            "row_number": self.row_number,
            "raw_row": self.raw_row,
            "error": self.error_message,
        }


class ExpenseProcessorResult:
    """Output of the expense processor — records plus any errors."""

    def __init__(
        self,
        records: list[ExpenseRecord],
        errors: list[ExpenseProcessingError],
    ):
        self.records = records
        self.errors = errors

    @property
    def success_count(self) -> int:
        return len(self.records)

    @property
    def error_count(self) -> int:
        return len(self.errors)

    def to_dict(self) -> dict:
        return {
            "records": [r.model_dump(mode="json") for r in self.records],
            "errors": [e.to_dict() for e in self.errors],
            "success_count": self.success_count,
            "error_count": self.error_count,
        }


def _parse_date(value: str) -> date:
    """Parse a date string in YYYY-MM-DD format."""
    parts = value.strip().split("-")
    return date(int(parts[0]), int(parts[1]), int(parts[2]))


def _validate_category(value: str) -> ExpenseCategory | None:
    """Return matching ExpenseCategory or None if unrecognized."""
    try:
        return ExpenseCategory(value)
    except ValueError:
        return None


def process_expenses(
    excel_path: Path | str,
    client_id: str,
    month: str,
) -> ExpenseProcessorResult:
    """Read an expense Excel file and return validated ExpenseRecord objects.

    Args:
        excel_path: Path to the .xlsx file.
        client_id: Client identifier to stamp on each record.
        month: Reporting month in YYYY-MM format.

    Returns:
        ExpenseProcessorResult with records and errors.
    """
    from openpyxl import load_workbook

    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"Expense file not found: {excel_path}")

    wb = load_workbook(str(excel_path), read_only=True, data_only=True)
    ws = wb.active

    records: list[ExpenseRecord] = []
    errors: list[ExpenseProcessingError] = []

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        wb.close()
        return ExpenseProcessorResult(records, errors)

    header = [str(h).strip() if h else "" for h in rows[0]]

    for row_idx, row_values in enumerate(rows[1:], start=2):
        raw: dict = {}
        for i, col_name in enumerate(EXPECTED_HEADERS):
            raw[col_name] = row_values[i] if i < len(row_values) else None

        try:
            employee = str(raw["Employee"] or "").strip()
            if not employee:
                raise ValueError("Missing employee name")

            date_str = str(raw["Date"] or "").strip()
            if not date_str:
                raise ValueError("Missing date")
            parsed_date = _parse_date(date_str)

            description = str(raw["Description"] or "").strip()
            if not description:
                raise ValueError("Missing description")

            vendor = str(raw["Vendor"] or "").strip()
            if not vendor:
                raise ValueError("Missing vendor")

            amount_raw = raw["Amount"]
            if amount_raw is None:
                raise ValueError("Missing amount")
            amount = float(amount_raw)
            if amount < 0:
                raise ValueError(f"Negative amount: {amount}")

            category_raw = str(raw["Category"] or "").strip()
            category = _validate_category(category_raw)
            if category is None:
                raise ValueError(f"Unrecognized category: {category_raw}")

            receipt_ref_raw = raw.get("Receipt Ref")
            receipt_ref = str(receipt_ref_raw).strip() if receipt_ref_raw else None
            if receipt_ref == "":
                receipt_ref = None

            record = ExpenseRecord(
                expense_id=f"EXP-{row_idx:04d}",
                client_id=client_id,
                employee=employee,
                date=parsed_date,
                description=description,
                vendor=vendor,
                amount=amount,
                category=category,
                receipt_reference=receipt_ref,
                status=ExpenseStatus.PENDING,
                month=month,
            )
            records.append(record)

        except (ValueError, TypeError) as exc:
            errors.append(ExpenseProcessingError(
                row_number=row_idx,
                raw_row=raw,
                error_message=str(exc),
            ))

    wb.close()
    return ExpenseProcessorResult(records, errors)
